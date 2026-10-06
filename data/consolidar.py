#!/usr/bin/env python3
"""
Consolida la base de la Notaría 2 de Cali en un solo archivo.

Fuentes (todas se unen en un esquema superset):
  - "Base de datos notaria 2.csv"  -> 1938-1943, transcripción vieja/parcial.
    Esquema pobre, valores en formato europeo ("3.360,00"), fecha = solo año.
    **1943 se EXCLUYE de esta fuente** porque fue re-transcrito completo en
    "1943.xlsx" (ver abajo). De aquí solo entran 1938, 1939, 1940 y 1941.
  - "1942.xlsx"                    -> 1942 (año que faltaba). Esquema rico.
  - "1943.xlsx"                    -> 1943 completo (~440 registros), esquema
    rico. REEMPLAZA los 66 registros viejos de 1943 (decisión del usuario,
    2026-07: el nuevo es la "tabla completa de transacciones").
  - "Base de datos 1944.xlsx"      -> 1944. Esquema rico. Precios en formato
    mixto US/europeo, fechas en texto libre ("11 de Enero de 1944").

Los tres archivos ricos (1942, 1943, 1944) comparten el mismo esquema de
columnas, pero 1942/1943 arrancan en la columna A y traen "N° Escritura",
mientras 1944 arranca en la columna C y no trae escritura.

Esquema de salida (superset):
  fuente, anio, fecha, vendedor, comprador, valor_texto, valor_num,
  area, negocio, escritura

Decisiones de diseño (acordadas con el usuario, 2026-06/07):
  - Superset: se conservan todas las columnas; los campos ausentes en una
    fuente (fecha completa y área en 1938-41; escritura en 1938-41 y 1944)
    quedan vacíos.
  - valor_num: columna numérica limpia; valor_texto conserva el original.
  - Partes (vendedor/comprador) se conservan tal cual; la limpieza de
    nombres y separación de apoderados se hará al armar el taller.
  - anio = año de la fuente (año de protocolización en la notaría), aunque
    algún acto proceda de una fecha anterior (p. ej. testamentos de 1940
    protocolizados en 1943).

Uso:
  python3 consolidar.py
"""
import csv
import re
import datetime
from pathlib import Path

import openpyxl

AQUI = Path(__file__).parent
VIEJO = AQUI / "Base de datos notaria 2.csv"
F1942 = AQUI / "1942.xlsx"
F1943 = AQUI / "1943.xlsx"
NUEVO = AQUI / "Base de datos 1944.xlsx"
SALIDA = AQUI / "consolidado_notaria2.csv"

# 1943 viejo se descarta: fue reemplazado por 1943.xlsx.
EXCLUIR_DEL_VIEJO = {"1943"}

CAMPOS = ["fuente", "anio", "fecha", "vendedor", "comprador",
          "valor_texto", "valor_num", "area", "negocio", "escritura"]


def limpiar_texto(s):
    """Quita espacios sobrantes y colapsa saltos de línea internos."""
    if s is None:
        return ""
    s = str(s).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def limpiar_escritura(v):
    """El N° de escritura llega de Excel como float (10.0). Se muestra como
    entero ('10') cuando es integral; los que son texto ('Test. Scarpetta')
    se limpian normal."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return limpiar_texto(v)


def parse_valor(s):
    """Convierte un precio en texto a float, detectando formato US vs europeo.

    Regla: si hay '.' y ',' el ÚLTIMO separador es el decimal.
    Si solo hay ',' decide por la cantidad de dígitos que le siguen
    (2 = decimal, 3 = miles). Si solo hay '.', misma regla (2 = decimal
    como "150.00"; 3 = miles como "9.000"). Rangos y guiones devuelven None.
    """
    if s is None:
        return None
    s = str(s).strip()
    if s in ("", "-", "No especificada"):
        return None
    s = s.replace("$", "").strip()
    # rangos tipo "380,00 - 2,780.80": ambiguo, no se normaliza
    if " - " in s or " – " in s:
        return None
    m = re.search(r"[\d.,]+", s)
    if not m:
        return None
    num = m.group().strip(".,")
    tiene_punto = "." in num
    tiene_coma = "," in num
    if tiene_punto and tiene_coma:
        if num.rfind(".") > num.rfind(","):   # US: 1,795.33
            num = num.replace(",", "")
        else:                                  # europeo: 1.500,00
            num = num.replace(".", "").replace(",", ".")
    elif tiene_coma:
        despues = num.split(",")[-1]
        if len(despues) == 2:                  # 175,75 -> decimal
            num = num.replace(",", ".")
        else:                                  # 40,000 -> miles
            num = num.replace(",", "")
    elif tiene_punto:
        despues = num.split(".")[-1]
        if len(despues) == 3:                  # 9.000 / 1.234.000 -> miles
            num = num.replace(".", "")
        # len == 2 ("150.00") u otros: float() lo maneja como decimal
    # sin separador: float() lo maneja directo
    try:
        return float(num)
    except ValueError:
        return None


def fecha_texto(valor):
    """Normaliza una fecha a texto legible.

    datetime -> 'YYYY-MM-DD'. String -> texto limpio ('-' se vuelve vacío).
    No fuerza el año: los acts con fecha anterior vienen como texto libre
    (p. ej. '28/10/1940 (prot. 1943)') y se conservan tal cual.
    """
    if valor is None:
        return ""
    if isinstance(valor, datetime.datetime):
        return valor.strftime("%Y-%m-%d")
    s = limpiar_texto(valor)
    return "" if s == "-" else s


def leer_viejo():
    """Lee la transcripción vieja (CSV), excluyendo los años ya reemplazados."""
    filas = []
    with open(VIEJO, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            anio = limpiar_texto(row.get("fecha"))
            if anio in EXCLUIR_DEL_VIEJO:
                continue
            filas.append({
                "fuente": "notaria2_1938-1941",
                "anio": anio,
                "fecha": "",
                "vendedor": limpiar_texto(row.get("vendedor")),
                "comprador": limpiar_texto(row.get("comprador")),
                "valor_texto": limpiar_texto(row.get("Valor")),
                "valor_num": parse_valor(row.get("Valor")),
                "area": "",
                "negocio": limpiar_texto(row.get("Negocio")),
                "escritura": "",
            })
    return filas


def leer_rico(path, anio, hoja="Hoja 1"):
    """Lee un Excel con el esquema rico A-G (1942, 1943).

    Columnas: A=N° Escritura, B=Fecha, C=Vendedor, D=Comprador,
    E=Precio, F=Área, G=Descripción. Fila 0 = título, fila 1 = encabezado.
    Se conserva toda fila que tenga vendedor o comprador.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[hoja]
    filas = []
    for i, row in enumerate(ws.iter_rows(min_col=1, max_col=7, values_only=True)):
        if i < 2:                              # título + encabezado
            continue
        escritura, fecha, vendedor, comprador, precio, area, desc = row
        vend = limpiar_texto(vendedor)
        comp = limpiar_texto(comprador)
        if not vend and not comp:              # fila vacía o de relleno
            continue
        filas.append({
            "fuente": f"notaria2_{anio}",
            "anio": str(anio),
            "fecha": fecha_texto(fecha),
            "vendedor": vend,
            "comprador": comp,
            "valor_texto": limpiar_texto(precio),
            "valor_num": parse_valor(precio),
            "area": limpiar_texto(area),
            "negocio": limpiar_texto(desc),
            "escritura": limpiar_escritura(escritura),
        })
    return filas


def leer_1944():
    """Lee el Excel de 1944 (esquema rico pero arranca en la columna C,
    sin columna de escritura y con fechas en texto libre)."""
    wb = openpyxl.load_workbook(NUEVO, data_only=True)
    ws = wb["Hoja1"]
    filas = []
    encabezado_visto = False
    for row in ws.iter_rows(min_col=3, max_col=8, values_only=True):
        if all(c is None for c in row):
            continue
        fecha, vendedor, comprador, precio, area, desc = row
        if not encabezado_visto:               # primera fila no vacía = encabezado
            encabezado_visto = True
            continue
        filas.append({
            "fuente": "notaria2_1944",
            "anio": "1944",
            "fecha": fecha_texto(fecha),
            "vendedor": limpiar_texto(vendedor),
            "comprador": limpiar_texto(comprador),
            "valor_texto": limpiar_texto(precio),
            "valor_num": parse_valor(precio),
            "area": limpiar_texto(area),
            "negocio": limpiar_texto(desc),
            "escritura": "",
        })
    return filas


def main():
    fuentes = [
        ("1938-1941 (viejo)", leer_viejo()),
        ("1942",             leer_rico(F1942, 1942)),
        ("1943",             leer_rico(F1943, 1943)),
        ("1944",             leer_1944()),
    ]
    todo = [fila for _, filas in fuentes for fila in filas]

    with open(SALIDA, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CAMPOS)
        w.writeheader()
        for fila in todo:
            if fila["valor_num"] is not None:
                fila = {**fila, "valor_num": f"{fila['valor_num']:.2f}"}
            else:
                fila = {**fila, "valor_num": ""}
            w.writerow(fila)

    # --- resumen ---
    for etiqueta, filas in fuentes:
        print(f"Registros {etiqueta:20}: {len(filas)}")
    print(f"Total consolidado:            {len(todo)}")
    sin_valor = sum(1 for r in todo if r["valor_num"] is None)
    print(f"Sin valor numérico:           {sin_valor} (rangos, guiones o vacíos)")
    print(f"Salida: {SALIDA.name}")


if __name__ == "__main__":
    main()
